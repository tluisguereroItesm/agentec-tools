from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path


def _load_input(path_arg: str) -> dict:
    return json.loads(Path(path_arg).read_text(encoding="utf-8"))


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _word_count(text: str) -> int:
    return len(text.split())


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


# ── Extractors ────────────────────────────────────────────────────────────────

def _extract_pdf(file_path: Path, max_chars: int) -> dict:
    try:
        import pdfplumber
    except ImportError:
        return {"success": False, "errorType": "MISSING_DEPENDENCY", "message": "pdfplumber no está instalado"}

    pages_text: list[str] = []
    with pdfplumber.open(str(file_path)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text() or ""
            pages_text.append(text)

    full_text = "\n".join(pages_text)
    content, truncated = _truncate(full_text, max_chars)
    return {
        "success": True,
        "content": content,
        "pageCount": page_count,
        "wordCount": _word_count(content),
        "charCount": len(content),
        "fileType": "pdf",
        "truncated": truncated,
    }


def _extract_docx(file_path: Path, max_chars: int) -> dict:
    try:
        from docx import Document
    except ImportError:
        return {"success": False, "errorType": "MISSING_DEPENDENCY", "message": "python-docx no está instalado"}

    doc = Document(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)
    content, truncated = _truncate(full_text, max_chars)
    return {
        "success": True,
        "content": content,
        "pageCount": None,
        "wordCount": _word_count(content),
        "charCount": len(content),
        "fileType": "docx",
        "truncated": truncated,
    }


def _extract_xlsx(file_path: Path, max_chars: int) -> dict:
    try:
        import openpyxl
    except ImportError:
        return {"success": False, "errorType": "MISSING_DEPENDENCY", "message": "openpyxl no está instalado"}

    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    rows: list[str] = []
    sheet_count = 0
    for sheet in wb.worksheets:
        sheet_count += 1
        rows.append(f"[Hoja: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append("\t".join(cells))
    wb.close()

    full_text = "\n".join(rows)
    content, truncated = _truncate(full_text, max_chars)
    return {
        "success": True,
        "content": content,
        "pageCount": sheet_count,
        "wordCount": _word_count(content),
        "charCount": len(content),
        "fileType": "xlsx",
        "truncated": truncated,
    }


def _extract_text(file_path: Path, max_chars: int) -> dict:
    full_text = file_path.read_text(encoding="utf-8", errors="replace")
    content, truncated = _truncate(full_text, max_chars)
    ext = file_path.suffix.lower().lstrip(".")
    return {
        "success": True,
        "content": content,
        "pageCount": None,
        "wordCount": _word_count(content),
        "charCount": len(content),
        "fileType": ext or "txt",
        "truncated": truncated,
    }


# ── Router ────────────────────────────────────────────────────────────────────

EXTRACTORS = {
    ".pdf":  _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".txt":  _extract_text,
    ".md":   _extract_text,
    ".csv":  _extract_text,
}


def extract(input_data: dict) -> dict:
    file_path = Path(input_data.get("filePath", ""))
    max_chars = int(input_data.get("maxChars", 8000))

    if not file_path.exists():
        return {
            "success": False,
            "message": f"Archivo no encontrado: {file_path}",
            "errorType": "FILE_NOT_FOUND",
        }

    ext = file_path.suffix.lower()
    extractor = EXTRACTORS.get(ext)

    if extractor is None:
        return {
            "success": False,
            "message": f"Formato no soportado: '{ext}'. Formatos válidos: {', '.join(EXTRACTORS)}",
            "errorType": "UNSUPPORTED_FORMAT",
        }

    try:
        result = extractor(file_path, max_chars)
    except Exception as exc:
        return {
            "success": False,
            "message": str(exc),
            "errorType": "EXTRACTION_ERROR",
        }

    result["message"] = (
        f"Texto extraído de {file_path.name}"
        + (" (truncado)" if result.get("truncated") else "")
    )
    result["extractedAt"] = _now_iso()
    return result


# ── Entrypoint ────────────────────────────────────────────────────────────────

def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "message": "Uso: python src/main.py <input.json>"}))
        sys.exit(1)

    input_data = _load_input(sys.argv[1])
    result = extract(input_data)

    # Persist artifact
    artifacts_dir = Path("/app/artifacts")
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
    artifact_path = artifacts_dir / f"doc-reader-{ts}.json"
    artifact_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    result["artifactPath"] = str(artifact_path)

    print(json.dumps(result, indent=2, ensure_ascii=False))
    sys.exit(0 if result.get("success") else 1)


if __name__ == "__main__":
    main()
